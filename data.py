"""Data loading and result saving for NCTracks Eligibility Verifier."""

import os
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from config import OUTPUT_COLUMNS, INPUT_COLUMNS


def generate_template(save_path: str) -> str:
    """Generate a blank patient list template Excel file."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Patient List"

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    header_align = Alignment(horizontal="center", wrap_text=True)

    for col_idx, header in enumerate(INPUT_COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    example = ["1234567890", "John", "Doe", "01/15/1990"]
    for col_idx, val in enumerate(example, 1):
        cell = ws.cell(row=2, column=col_idx, value=val)
        cell.font = Font(italic=True, color="888888")

    for col in ws.columns:
        col_letter = col[0].column_letter
        max_len = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    wb.save(save_path)
    return save_path


def load_patients(file_path: str) -> list[dict]:
    """Load patient list from Excel (.xlsx) or CSV (.csv) file."""
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".csv":
        return _load_csv(file_path)
    elif ext in (".xlsx", ".xls"):
        return _load_excel(file_path)
    else:
        raise ValueError(f"Unsupported file format: {ext}. Use .csv or .xlsx")


def _load_csv(file_path: str) -> list[dict]:
    import csv

    patients = []
    with open(file_path, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        _validate_columns(reader.fieldnames, file_path)
        for row in reader:
            patients.append(_normalize_row(row))
    return patients


def _load_excel(file_path: str) -> list[dict]:
    wb = openpyxl.load_workbook(file_path, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError(f"Empty spreadsheet: {file_path}")

    headers = [str(h).strip() if h else "" for h in rows[0]]
    _validate_columns(headers, file_path)

    patients = []
    for row in rows[1:]:
        if all(cell is None for cell in row):
            continue
        row_dict = dict(zip(headers, row))
        patients.append(_normalize_row(row_dict))

    wb.close()
    return patients


def _validate_columns(headers: list[str] | None, file_path: str):
    if headers is None:
        raise ValueError(f"Could not read headers from {file_path}")
    headers_upper = [h.upper().strip() for h in headers]
    for col in INPUT_COLUMNS:
        if col.upper() not in headers_upper:
            raise ValueError(
                f"Missing required column '{col}' in {file_path}. "
                f"Expected columns: {INPUT_COLUMNS}"
            )


def _normalize_row(row: dict) -> dict:
    """Normalize a row dict to standard keys."""
    normalized = {}
    for key, value in row.items():
        if key is None:
            continue
        k = key.strip().upper()
        val = str(value).strip() if value is not None else ""
        if k == "MEDICAID ID":
            normalized["medicaid_id"] = val
        elif k == "FIRST NAME":
            normalized["first_name"] = val
        elif k == "LAST NAME":
            normalized["last_name"] = val
        elif k == "DOB":
            normalized["dob"] = val
    return normalized


# ─── Status colors ───
_STATUS_FILLS = {
    "ELIGIBLE": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
    "NOT ELIGIBLE": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
    "NOT FOUND": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
    "MANAGED CARE": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
    "FFS": PatternFill(start_color="FFE599", end_color="FFE599", fill_type="solid"),
    "ERROR": PatternFill(start_color="F4CCCC", end_color="F4CCCC", fill_type="solid"),
    "SKIPPED": PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid"),
}
_PAYER_CHANGED_FONT = Font(bold=True, color="CC0000")


def save_results(results: list[dict], output_path: str):
    """Save verification results to a styled Excel file."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Eligibility Results"

    # Header styling
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    header_align = Alignment(horizontal="center", wrap_text=True)

    for col_idx, header in enumerate(OUTPUT_COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    # Write data rows
    for row_idx, r in enumerate(results, 2):
        emr_source = r.get("emr_funding_source", "")
        entity = r.get("managing_entity", "")
        entity_next = r.get("managing_entity_next", "")
        no_entity = not entity or entity == "(none)"
        no_entity_next = not entity_next or entity_next == "(none)"

        row_data = [
            r.get("medicaid_id", ""),
            r.get("name", ""),
            r.get("status", ""),
            emr_source,
            r.get("insurance_type", ""),
            "None" if no_entity else entity,
            r.get("current_period", ""),
            "None" if no_entity else r.get("payer_changed", ""),
            "None" if no_entity_next else entity_next,
            r.get("next_period", ""),
            "None" if no_entity_next else r.get("payer_changed_next", ""),
            r.get("checked_at", ""),
            r.get("notes", ""),
        ]

        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)

        # Status color (column 3)
        status = r.get("status", "")
        if status in _STATUS_FILLS:
            ws.cell(row=row_idx, column=3).fill = _STATUS_FILLS[status]

        # Payer Changed highlighting (columns 8 and 11)
        for col in [8, 11]:
            cell = ws.cell(row=row_idx, column=col)
            if cell.value == "YES":
                cell.font = _PAYER_CHANGED_FONT

    # Summary row
    summary_row = len(results) + 3
    eligible = sum(1 for r in results if r.get("status") in ("ELIGIBLE", "MANAGED CARE"))
    not_eligible = sum(1 for r in results if r.get("status") in ("NOT ELIGIBLE", "NOT FOUND"))
    errors = sum(1 for r in results if r.get("status") == "ERROR")
    skipped = sum(1 for r in results if r.get("status") == "SKIPPED")

    ws.cell(row=summary_row, column=1, value="Summary").font = Font(bold=True)
    ws.cell(row=summary_row + 1, column=1, value=f"Eligible: {eligible}")
    ws.cell(row=summary_row + 2, column=1, value=f"Not Eligible: {not_eligible}")
    ws.cell(row=summary_row + 3, column=1, value=f"Errors: {errors}")
    if skipped > 0:
        ws.cell(row=summary_row + 4, column=1, value=f"Skipped: {skipped}")

    # Column widths
    col_widths = [14, 25, 14, 30, 14, 28, 16, 16, 28, 16, 18, 20, 40]
    for i, width in enumerate(col_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width

    wb.save(output_path)
    return output_path


def generate_output_path(input_path: str | None = None) -> str:
    """Generate output file path based on input file name or timestamp."""
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    if input_path:
        base = os.path.splitext(os.path.basename(input_path))[0]
        directory = os.path.dirname(input_path) or "."
        return os.path.join(directory, f"{base}_results_{timestamp}.xlsx")
    else:
        return f"eligibility_results_{timestamp}.xlsx"
